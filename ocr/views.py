#all bugs
#document and pdf not download encoding error in some images
#show error if found error in extracting the text
#all text dosen't come in pdf file


from django.shortcuts import render, redirect
from .forms import UserFileForm, FileModelForm
from .models import FileModel,ContactUs
from .ocr_engine import get_text
from django.contrib.auth.models import User
from django.contrib.auth import authenticate,  login, logout
from django.contrib.auth.forms import UserChangeForm
from django.http import HttpResponse
from django.contrib import messages
from wsgiref.util import FileWrapper
import os
from fpdf import FPDF
from django.http import FileResponse
from docx import Document
import pytesseract as pt
import PIL
from PIL import Image
pt.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
from google_trans_new import google_translator
from django.utils.datastructures import MultiValueDictKeyError

LANGUAGES = {
    'af': 'afrikaans',  'sq': 'albanian',    'am': 'amharic',    'ar': 'arabic',    'hy': 'armenian',    'az': 'azerbaijani',    'eu': 'basque',
    'be': 'belarusian',    'bn': 'bengali',    'bs': 'bosnian',    'bg': 'bulgarian',    'ca': 'catalan',    'ceb': 'cebuano',    'ny': 'chichewa',
    'zh-cn': 'chinese (simplified)',    'zh-tw': 'chinese (traditional)',    'co': 'corsican',    'hr': 'croatian',    'cs': 'czech',    'da': 'danish',
    'nl': 'dutch',    'en': 'english',    'eo': 'esperanto',    'et': 'estonian',    'tl': 'filipino',    'fi': 'finnish',    'fr': 'french',
    'fy': 'frisian',    'gl': 'galician',    'ka': 'georgian',    'de': 'german',    'el': 'greek',    'gu': 'gujarati',    'ht': 'haitian creole',
    'ha': 'hausa',    'haw': 'hawaiian',    'iw': 'hebrew',    'he': 'hebrew',    'hi': 'hindi',    'hmn': 'hmong',    'hu': 'hungarian',
    'is': 'icelandic',    'ig': 'igbo',    'id': 'indonesian',    'ga': 'irish',    'it': 'italian',    'ja': 'japanese',    'jw': 'javanese',
    'kn': 'kannada',    'kk': 'kazakh',    'km': 'khmer',    'ko': 'korean',    'ku': 'kurdish (kurmanji)',    'ky': 'kyrgyz',    'lo': 'lao',
    'la': 'latin',    'lv': 'latvian',    'lt': 'lithuanian',    'lb': 'luxembourgish',    'mk': 'macedonian',    'mg': 'malagasy',    'ms': 'malay',
    'ml': 'malayalam',    'mt': 'maltese',    'mi': 'maori',    'mr': 'marathi',    'mn': 'mongolian',    'my': 'myanmar (burmese)',    'ne': 'nepali',
    'no': 'norwegian',    'or': 'odia',    'ps': 'pashto',    'fa': 'persian',    'pl': 'polish',    'pt': 'portuguese',    'pa': 'punjabi',    'ro': 'romanian',
    'ru': 'russian',    'sm': 'samoan',    'gd': 'scots gaelic',    'sr': 'serbian',    'st': 'sesotho',    'sn': 'shona',    'sd': 'sindhi',
    'si': 'sinhala',    'sk': 'slovak',    'sl': 'slovenian',    'so': 'somali',    'es': 'spanish',    'su': 'sundanese',    'sw': 'swahili',
    'sv': 'swedish',    'tg': 'tajik',    'ta': 'tamil',    'te': 'telugu',    'th': 'thai',    'tr': 'turkish',    'uk': 'ukrainian',    'ur': 'urdu',
    'ug': 'uyghur',    'uz': 'uzbek',    'vi': 'vietnamese',    'cy': 'welsh',    'xh': 'xhosa',    'yi': 'yiddish',    'yo': 'yoruba',    'zu': 'zulu',
}

#get text and serve to download in 3 file formates


def download_pdf(request):
   file_data = request.POST["txtareaField"]
   if request.POST.get("txt"):
       response = HttpResponse(file_data, content_type='application/text charset=utf-8')
       response['Content-Disposition'] = 'attachment; filename="download.txt"'
       return response

   if request.POST.get("pdf"):
       if os.path.exists('ocr/static/ocr/img/download.pdf'):
        os.remove("ocr/static/ocr/img/download.pdf")
       document=FPDF()
       document.add_page()
       document.set_font("Arial", size=15)
       document.cell(200, 10, txt=file_data, ln=1, align="L")
       document.output("ocr/static/ocr/img/download.pdf")
       document=FPDF(orientation='P', unit='mm', format='A3')
       print("pdf has been created successfully....")
       file = open('ocr/static/ocr/img/download.pdf', 'rb')
       response = FileResponse(file)
       response['Content-Type'] = 'application/application/pdf'
       response['Content-Disposition'] = 'attachment;filename=download.pdf'  # This file name is the default file name when downloading and can be changed

       #this will open file in new tab to read
       #response = FileResponse(open('ocr/static/ocr/img/download.pdf', 'rb'))
       return response

   if request.POST.get("docx"):
       if os.path.exists('ocr/static/ocr/img/download.docx'):
        os.remove("ocr/static/ocr/img/download.docx")
       document = Document()
       document.add_paragraph(file_data, style='IntenseQuote')
       document.save('ocr/static/ocr/img/download.docx')
       response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
       response['Content-Disposition'] = 'attachment; filename=download.docx'
       document.save(response)
       return response


def homepage(request):
    bool = "false"
    if request.user.is_authenticated:
        bool="true"

    context = {'form': UserFileForm(),'bool':bool,"L":LANGUAGES }
    if request.method=='POST' and UserFileForm(request.POST, request.FILES).is_valid():
        user_input = UserFileForm(request.POST, request.FILES)
        if user_input.is_valid():
            fileobj = request.FILES['userfile']
            lang = user_input.cleaned_data['lang']
            is_image = fileobj.name.endswith(('jpg', 'png', 'jpeg', 'tiff'))
            if is_image:
                multi = False
            else:
                multi = True
            #text = get_text(fileobj, lang, multi)

            text = pt.image_to_string(Image.open(fileobj), lang='eng')
            context['text'] = text
            return render(request, 'ocr/index.html', context)
        else:
            context['form_error'] = 'Error! File can not be converted for recognition'
            return render(request, 'ocr/index.html', context)
    elif request.method=='POST':
        file_data = request.POST["txtareaField"]
        Language = request.POST.get("lang", None)
        translator = google_translator()  
        translate_text = translator.translate(file_data,lang_tgt=Language)
        context['text'] = file_data
        context['transletedText'] =  translate_text
        # print(translate_text)

    return render(request, 'ocr/index.html', context)


# handel user login reuest

def userLogin(request):
    if request.method == "POST":
        loginusername = request.POST['loginusername']
        loginpassword = request.POST['loginpassword']

        user = authenticate(username=loginusername, password=loginpassword)

        if user is not None:
            login(request, user)
            return redirect("homepage")
            messages.info(request, 'You are logged in succesffully!')
        else:
             messages.info(request, 'Username or password is invalid, Please check it once.!')
             return render(request, 'ocr/error404.html')

    else:
        return render(request, 'ocr/error404.html')


def userLogout(request):
    logout(request)
    return redirect('homepage')


def userSignup(request):
    if request.method == "POST":
        username = request.POST['username']
        email = request.POST['email']
        fname = request.POST['fname']
        lname = request.POST['lname']
        pass1 = request.POST['pass1']
        pass2 = request.POST['pass2']

        # check for errorneous input

        if (pass1 != pass2):
            messages.info(request, 'Your passsword is not matching')
            return redirect('homepage')

        # chekc username is unique or not and Create new user
        try:
            user= User.objects.get(username=username)
            messages.info(request, 'Your username is not unique. try another one')
            return redirect('homepage')
        except User.DoesNotExist:
            myuser = User.objects.create_user(username, email, pass1)
            myuser.first_name = fname
            myuser.last_name = lname
            myuser.save()
            login(request, myuser)
            messages.info(request, 'Welcome to Django-ocr app! You have created your account successfully')
            return redirect('homepage')
    else:
        return render(request, 'ocr/error404.html')

#contact us model
def Contact(request):
    if request.method == 'POST':
        form = ContactUs()
        name = request.POST.get('name')
        email = request.POST.get('email')
        message = request.POST.get('message')

        form.name = name
        form.mail = email
        form.message = message
        form.save()
        messages.info(request, 'We got your query. we will look at soon.')
        return render(request,'ocr/contactus.html')
    return render(request,'ocr/contactus.html')
